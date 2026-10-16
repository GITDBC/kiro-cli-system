"""Document conversion through MarkItDown in a process pool."""

# Standard libraries
import logging
import os
import re
import shutil
from collections import defaultdict
from collections.abc import Iterator
from concurrent.futures import Future, ProcessPoolExecutor, as_completed
from pathlib import Path
from typing import Any
from urllib.parse import quote

# 3pps
from markitdown import MarkItDown

# Own modules
from doc2md.types import (
    ALLOWED_IMG_EXTS,
    IMAGE_EXTS,
    MAX_IMAGE_BLOB_SIZE,
    PPTX_EXTS,
    ConvertResult,
    ResultCallback,
)

logger: logging.Logger = logging.getLogger(__name__)

_converter: MarkItDown | None = None


def _init_worker() -> None:
    """
    Initialise a per-process MarkItDown instance.

    Called once by each worker in the process pool to
    create a module-level converter that is reused for
    every file the worker handles.

    Returns:
        None.
    """

    global _converter
    _converter = MarkItDown()


def _iter_shapes(shapes: Any) -> Iterator[Any]:
    """
    Recursively yield all shapes, descending into groups.

    Args:
        shapes: Top-level shape collection from a slide.

    Yields:
        Individual shapes, including those nested in
            groups.
    """

    # 3pps
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for s in shapes:
        yield s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(s.shapes)


def _process_shape(
    shape: Any,
    slide_idx: int,
    img_dir: Path,
    file_stem: str,
    replacements: dict[str, list[str]],
) -> None:
    """
    Extract a single shape's image blob to disk and
    record the fake-to-real name replacement.

    Args:
        shape: A PowerPoint shape object.
        slide_idx: Zero-based slide index.
        img_dir: Directory for extracted images.
        file_stem: Stem of the source PPTX file.
        replacements: Mutable mapping of fake names to
            real URL-encoded references.

    Returns:
        None.
    """

    if not hasattr(shape, "image"):
        return

    try:
        blob: bytes = shape.image.blob
    except Exception:
        logger.debug("Cannot read image from shape '%s'.", shape.name, exc_info=True)
        return

    if len(blob) > MAX_IMAGE_BLOB_SIZE:
        logger.warning("Oversized image blob in '%s', skipping.", shape.name)
        return

    sanitized: str = re.sub(r"\W", "", shape.name) or f"shape{slide_idx}_{id(shape)}"
    fake_name: str = f"{sanitized}.jpg"
    img_filename: str = shape.image.filename or ""
    raw_ext: str = Path(img_filename).suffix.lower()
    real_ext: str = raw_ext if raw_ext in ALLOWED_IMG_EXTS else ".png"
    real_name: str = f"s{slide_idx}_{sanitized}{real_ext}"

    (img_dir / real_name).write_bytes(blob)

    ref: str = quote(f"{file_stem}_images/{real_name}")
    replacements[fake_name].append(ref)


def _apply_image_replacements(
    content: str,
    replacements: dict[str, list[str]],
) -> str:
    """
    Replace fake image references in Markdown content
    with real extracted image paths.

    Args:
        content: Markdown content with placeholder refs.
        replacements: Mapping of fake filenames to lists
            of real URL-encoded references.

    Returns:
        Markdown content with corrected image paths.
    """

    for fake, refs in replacements.items():
        ref_iter = iter(refs)

        def _sub(m: re.Match[str], it: Iterator[str] = ref_iter) -> str:
            ref = next(it, None)
            return f"]({ref})" if ref else m.group(0)

        content = re.sub(
            re.escape(f"]({fake})"),
            _sub,
            content,
            count=len(refs),
        )

    return content


def _extract_pptx_images(
    file: Path,
    dest_dir: Path,
    content: str,
) -> str:
    """
    Extract embedded images from a PowerPoint file and
    rewrite the Markdown image references to point to the
    extracted files on disk.

    MarkItDown generates placeholder filenames for images
    that do not match the actual embedded format. This
    function corrects those references and saves the real
    image blobs.

    Args:
        file: Source ``.pptx`` file.
        dest_dir: Directory where the Markdown file will
            be written.
        content: Markdown content produced by MarkItDown
            with placeholder image references.

    Returns:
        Markdown content with corrected image paths.
    """

    # 3pps
    import pptx

    img_dir: Path = dest_dir / f"{file.stem}_images"
    if not img_dir.resolve().is_relative_to(dest_dir.resolve()):
        return content
    img_dir.mkdir(parents=True, exist_ok=True)

    replacements: dict[str, list[str]] = defaultdict(list)

    for slide_idx, slide in enumerate(pptx.Presentation(str(file)).slides):
        for shape in _iter_shapes(slide.shapes):
            _process_shape(shape, slide_idx, img_dir, file.stem, replacements)

    return _apply_image_replacements(content, replacements)


def _convert_single(
    file: Path,
    output_dir: Path,
    rel_to: Path | None = None,
) -> ConvertResult:
    """
    Convert a single non-audio file to Markdown using
    MarkItDown.

    Images that produce no text content are copied as-is
    to the output directory. PowerPoint files receive
    additional post-processing to extract embedded images.

    Args:
        file: Source file to convert.
        output_dir: Root directory for generated Markdown.
        rel_to: Common ancestor path used to preserve
            directory structure in the output.

    Returns:
        Tuple of (source path, destination path or None
            on failure, error message or None on success).
    """

    try:
        if _converter is None:
            return file, None, "MarkItDown converter not initialised."

        content: str = _converter.convert(str(file)).text_content or ""

        relative: Path = file.relative_to(rel_to) if rel_to else Path(file.name)
        dest_dir: Path = output_dir / relative.parent

        if not dest_dir.resolve().is_relative_to(output_dir.resolve()):
            return file, None, "output path escapes output directory."

        dest_dir.mkdir(parents=True, exist_ok=True)

        ext: str = file.suffix.lower()

        if not content.strip() and ext in IMAGE_EXTS:
            shutil.copy2(file, dest_dir / file.name, follow_symlinks=False)
            return file, dest_dir / file.name, None

        if not content.strip():
            return file, None, "conversion produced empty output."

        if ext in PPTX_EXTS:
            content = _extract_pptx_images(file, dest_dir, content)

        dest: Path = dest_dir / f"{file.stem}_{file.suffix.lstrip('.')}.md"
        dest.write_text(content, encoding="utf-8")

    except Exception as e:
        logger.debug("Conversion failed for '%s'.", file, exc_info=True)
        return file, None, str(e)

    else:
        return file, dest, None


def convert_documents(
    files: list[Path],
    output_dir: Path,
    rel_to: Path | None,
    on_result: ResultCallback,
) -> int:
    """
    Convert document files to Markdown in parallel using
    a process pool.

    Each worker process holds its own MarkItDown instance.
    Results are reported through the callback as they
    complete.

    Args:
        files: Document files to convert.
        output_dir: Root directory for generated Markdown.
        rel_to: Common ancestor path used to preserve
            directory structure in the output.
        on_result: Callback invoked with each conversion
            result as ``(source, dest, error)``. Must
            return a truthy value on success.

    Returns:
        Number of files successfully converted.
    """

    converted: int = 0

    with ProcessPoolExecutor(
        max_workers=min(os.cpu_count() or 1, 4),
        initializer=_init_worker,
    ) as pool:
        futures: dict[Future[ConvertResult], Path] = {
            pool.submit(_convert_single, f, output_dir, rel_to): f for f in files
        }

        for future in as_completed(futures):
            try:
                converted += on_result(*future.result())
            except Exception as e:
                src = futures[future]
                converted += on_result(src, None, str(e))

    return converted
